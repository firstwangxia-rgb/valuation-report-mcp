"""Filesystem-safe core logic shared by MCP tools and tests."""

from __future__ import annotations

import hashlib
import os
import re
import tempfile
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from valuation_report_mcp.schemas import (
    CompanyKnowledgeRequest,
    ComplianceItem,
    ComplianceResult,
    DraftReportRequest,
    KnowledgeCategory,
    KnowledgeInventory,
    ResearchPlan,
    SearchHit,
    SearchResponse,
    SourceLink,
    WriteResult,
)

KNOWLEDGE_CATEGORIES = ("准则", "评估报告", "评估说明", "公司档案")
SEARCHABLE_SUFFIXES = {".md", ".txt", ".pdf", ".docx"}
DEFAULT_MAX_FILE_BYTES = 25 * 1024 * 1024
DEFAULT_MAX_EXTRACTED_CHARS = 200_000

REQUIRED_REPORT_SECTIONS = (
    "报告声明",
    "摘要",
    "委托方",
    "被评估单位",
    "评估目的",
    "评估对象",
    "评估基准日",
    "评估方法",
    "价值类型",
    "评估假设",
    "评估结论",
    "特别事项",
)
ALLOWED_METHODS = ("收益法", "市场法", "资产基础法", "成本法")
WINDOWS_RESERVED_NAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


@dataclass(frozen=True)
class WorkspaceConfig:
    root: Path
    knowledge_base: Path
    reports: Path
    max_file_bytes: int = DEFAULT_MAX_FILE_BYTES
    max_extracted_chars: int = DEFAULT_MAX_EXTRACTED_CHARS

    @classmethod
    def from_env(cls) -> WorkspaceConfig:
        root = Path(os.getenv("VALUATION_MCP_WORKSPACE_ROOT", os.getcwd())).expanduser().resolve()
        knowledge_base = resolve_within(
            root, os.getenv("VALUATION_MCP_KNOWLEDGE_DIR", "knowledge-base")
        )
        reports = resolve_within(root, os.getenv("VALUATION_MCP_REPORTS_DIR", "reports"))
        max_file_bytes = _positive_int_env("VALUATION_MCP_MAX_FILE_BYTES", DEFAULT_MAX_FILE_BYTES)
        max_extracted_chars = _positive_int_env(
            "VALUATION_MCP_MAX_EXTRACTED_CHARS", DEFAULT_MAX_EXTRACTED_CHARS
        )
        return cls(root, knowledge_base, reports, max_file_bytes, max_extracted_chars)


def _positive_int_env(name: str, default: int) -> int:
    raw = os.getenv(name)
    if raw is None:
        return default
    value = int(raw)
    if value <= 0:
        raise ValueError(f"{name} must be a positive integer")
    return value


def resolve_within(root: Path, relative_path: str | Path) -> Path:
    """Resolve a relative path and reject escapes outside root."""

    candidate_path = Path(relative_path)
    if candidate_path.is_absolute():
        raise ValueError("absolute paths are not allowed")
    root = root.resolve()
    candidate = (root / candidate_path).resolve()
    if candidate != root and root not in candidate.parents:
        raise ValueError("path escapes the configured workspace root")
    return candidate


def sanitize_markdown_filename(value: str) -> str:
    """Convert a user-facing title into a safe single Markdown filename."""

    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")
    cleaned = re.sub(r"\s+", " ", cleaned)
    if not cleaned:
        raise ValueError("filename is empty after sanitization")
    stem = Path(cleaned).stem
    if stem.upper() in WINDOWS_RESERVED_NAMES:
        cleaned = f"_{cleaned}"
    if not cleaned.lower().endswith(".md"):
        cleaned += ".md"
    if len(cleaned) > 160:
        suffix = ".md"
        cleaned = cleaned[: 160 - len(suffix)].rstrip(" .") + suffix
    return cleaned


def _read_text_file(path: Path, max_chars: int) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)[:max_chars]
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")[:max_chars]


def extract_text(path: Path, max_chars: int) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return _read_text_file(path, max_chars)
    if suffix == ".pdf":
        parts: list[str] = []
        total = 0
        for page in PdfReader(str(path)).pages:
            text = page.extract_text() or ""
            parts.append(text)
            total += len(text)
            if total >= max_chars:
                break
        return "\n".join(parts)[:max_chars]
    if suffix == ".docx":
        parts = []
        total = 0
        document = Document(str(path))
        for paragraph in document.paragraphs:
            parts.append(paragraph.text)
            total += len(paragraph.text)
            if total >= max_chars:
                break
        if total < max_chars:
            for table in document.tables:
                for row in table.rows:
                    row_text = "\t".join(cell.text for cell in row.cells)
                    parts.append(row_text)
                    total += len(row_text)
                    if total >= max_chars:
                        break
                if total >= max_chars:
                    break
        return "\n".join(parts)[:max_chars]
    raise ValueError(f"unsupported file type: {suffix}")


def _iter_searchable_files(config: WorkspaceConfig, categories: Iterable[str]) -> Iterable[Path]:
    for category in categories:
        category_dir = resolve_within(config.knowledge_base, category)
        if not category_dir.exists():
            continue
        for path in category_dir.rglob("*"):
            if (
                path.is_file()
                and not path.is_symlink()
                and path.suffix.lower() in SEARCHABLE_SUFFIXES
                and not any(part.startswith(".") for part in path.relative_to(category_dir).parts)
            ):
                yield path


def _make_snippet(text: str, terms: list[str], radius: int = 180) -> str:
    lowered = text.casefold()
    positions = [lowered.find(term.casefold()) for term in terms]
    positions = [position for position in positions if position >= 0]
    if not positions:
        return ""
    start = max(0, min(positions) - radius)
    end = min(len(text), start + radius * 2)
    snippet = re.sub(r"\s+", " ", text[start:end]).strip()
    return ("…" if start else "") + snippet + ("…" if end < len(text) else "")


def search_knowledge(
    config: WorkspaceConfig,
    query: str,
    keywords: list[str] | None = None,
    categories: list[str] | None = None,
    limit: int = 20,
) -> SearchResponse:
    query = query.strip()
    if not query:
        raise ValueError("query must not be empty")
    if not 1 <= limit <= 50:
        raise ValueError("limit must be between 1 and 50")

    selected_categories = categories or list(KNOWLEDGE_CATEGORIES)
    invalid_categories = sorted(set(selected_categories) - set(KNOWLEDGE_CATEGORIES))
    if invalid_categories:
        raise ValueError(f"unsupported categories: {', '.join(invalid_categories)}")

    terms = list(dict.fromkeys([query, *(keywords or [])]))
    terms = [term.strip() for term in terms if term.strip()]
    lowered_terms = [term.casefold() for term in terms]
    hits: list[SearchHit] = []
    warnings: list[str] = []
    searched_files = 0
    skipped_files = 0

    for path in _iter_searchable_files(config, selected_categories):
        try:
            size = path.stat().st_size
        except OSError as exc:
            skipped_files += 1
            warnings.append(f"无法读取文件信息 {path.name}: {exc}")
            continue
        if size > config.max_file_bytes:
            skipped_files += 1
            warnings.append(f"文件过大，已跳过：{path.name}")
            continue

        searched_files += 1
        try:
            text = extract_text(path, config.max_extracted_chars)
        except Exception as exc:  # extraction libraries raise several format-specific exceptions
            skipped_files += 1
            warnings.append(f"提取失败 {path.name}: {type(exc).__name__}")
            continue

        name_folded = path.name.casefold()
        text_folded = text.casefold()
        matched_terms = [
            term
            for term, lowered in zip(terms, lowered_terms, strict=True)
            if lowered in name_folded or lowered in text_folded
        ]
        if not matched_terms:
            continue

        score = sum(
            name_folded.count(lowered) * 20 + min(text_folded.count(lowered), 20)
            for lowered in lowered_terms
        )
        relative = path.relative_to(config.root)
        category = relative.parts[1] if len(relative.parts) > 1 else ""
        hits.append(
            SearchHit(
                path=relative.as_posix(),
                category=category,
                file_type=path.suffix.lower().lstrip("."),
                score=score,
                matched_terms=matched_terms,
                snippet=_make_snippet(text, matched_terms),
            )
        )

    hits.sort(key=lambda hit: (-hit.score, hit.path))
    return SearchResponse(
        query=query,
        workspace_root=str(config.root),
        searched_files=searched_files,
        skipped_files=skipped_files,
        hits=hits[:limit],
        warnings=warnings[:50],
    )


def list_knowledge(config: WorkspaceConfig) -> KnowledgeInventory:
    categories: list[KnowledgeCategory] = []
    for name in KNOWLEDGE_CATEGORIES:
        path = resolve_within(config.knowledge_base, name)
        files = (
            [
                item
                for item in path.rglob("*")
                if path.exists() and item.is_file() and not item.is_symlink()
            ]
            if path.exists()
            else []
        )
        categories.append(
            KnowledgeCategory(
                name=name,
                path=path.relative_to(config.root).as_posix(),
                file_count=len(files),
                total_bytes=sum(item.stat().st_size for item in files),
            )
        )
    return KnowledgeInventory(workspace_root=str(config.root), categories=categories)


def prepare_research_plan(
    company_name: str, extra_keywords: list[str] | None = None
) -> ResearchPlan:
    company_name = company_name.strip()
    if not company_name:
        raise ValueError("company_name must not be empty")
    queries = [
        f"{company_name} 企业概况 主营业务 官方",
        f"{company_name} 最近一年 年度报告 营业收入 净利润",
        f"{company_name} 股权结构 实际控制人 重大事项",
        f"{company_name} 行业地位 可比公司 竞争格局",
    ]
    for keyword in extra_keywords or []:
        keyword = keyword.strip()
        if keyword:
            queries.append(f"{company_name} {keyword}")
    return ResearchPlan(
        company_name=company_name,
        local_check_first=True,
        queries=list(dict.fromkeys(queries)),
        preferred_sources=[
            "证券交易所、监管机构及政府网站",
            "公司官网、年度报告、招股说明书和正式公告",
            "行业协会和权威统计机构",
            "媒体报道仅作辅助，并回查一手来源",
        ],
        required_fields=[
            "公司名称、成立时间、注册地址、实际控制人",
            "主营业务、产品结构、行业分类和竞争地位",
            "最近三年及最新一期财务数据",
            "重大融资、诉讼、监管、关联交易和期后事项",
            "每项关键事实对应的标题、发布日期和URL",
        ],
        host_instruction=(
            "请由具备联网能力的MCP宿主执行上述查询。优先打开一手来源，"
            "区分已审计、已审阅、公司预计和媒体转述数据；将结构化结果和URL"
            "传给 draft_valuation_report 或 update_company_knowledge。"
        ),
    )


def _source_markdown(sources: list[SourceLink]) -> str:
    if not sources:
        return "- 来源待补充；正式使用前必须补齐并复核。"
    return "\n".join(f"- [{source.title}]({source.url})" for source in sources)


def _bullet_markdown(items: list[str], fallback: str) -> str:
    return "\n".join(f"- {item}" for item in items) if items else f"- {fallback}"


def render_report(request: DraftReportRequest) -> str:
    disclaimer = (
        "本报告由AI辅助生成，仅用于教学、研究和正式评估前期讨论。未经现场调查、"
        "权属核验、底稿复核、资产评估师签名和评估机构盖章，不构成正式资产评估报告或投资建议。"
    )
    assumptions = _bullet_markdown(request.assumptions, "相关假设待委托方和评估师确认。")
    special_matters = _bullet_markdown(
        request.special_matters, "未提供特别事项；正式出具前应重新核验。"
    )
    manual_review = _bullet_markdown(
        request.manual_review_items,
        "委托方、评估目的、基准日、申报范围、评估参数和最终结论均需人工复核。",
    )
    sources = _source_markdown(request.sources)
    return f"""# {request.company_name}企业价值评估报告（初稿）

> {disclaimer}

## 一、报告声明

本报告根据用户提供的资料、本地知识库和列明的公开来源编制。所有事实、参数和结论均应由报告使用人及资产评估师复核。评估结论不等同于实际交易价格。

准则依据：请核对本地知识库中现行有效的《资产评估执业准则——企业价值》和《资产评估执业准则——资产评估报告》。

## 二、摘要

- 被评估单位：{request.company_name}
- 评估目的：{request.valuation_purpose}
- 评估基准日：{request.valuation_date}
- 价值类型：{request.value_type}
- 初步结论：{request.preliminary_conclusion}

## 三、委托方及被评估单位概况

### （一）委托方

委托方资料待正式业务承接时补充。

### （二）被评估单位

{request.company_profile}

## 四、评估目的与评估对象

评估目的：{request.valuation_purpose}

评估对象：{request.company_name}股东全部权益价值。正式项目应根据经济行为文件确认评估对象和评估范围。

## 五、评估基准日

评估基准日为{request.valuation_date}。基准日应与经济行为、财务资料和市场参数时点保持一致。

## 六、财务状况与经营成果

{request.financial_data}

## 七、行业状况与可比案例

{request.industry_comparison}

## 八、评估方法与价值类型

价值类型为{request.value_type}。

{request.method_analysis}

正式评估应分别分析收益法、市场法和资产基础法的适用性；适合采用多种方法时，应采用两种以上方法并说明最终结论的形成过程。

## 九、评估假设与限制条件

{assumptions}

## 十、评估结论

{request.preliminary_conclusion}

该结论为资料受限条件下的初步判断。缺少关键资料或程序时，不应表述为无保留的正式评估结论。

## 十一、特别事项说明

{special_matters}

## 十二、需人工复核事项

{manual_review}

## 十三、资料来源

{sources}

---

生成状态：Markdown初稿
生成工具：valuation-report-mcp
"""


def _atomic_write(target: Path, content: str, overwrite: bool) -> WriteResult:
    if target.exists() and target.is_symlink():
        raise ValueError("refusing to write through a symlink")
    target.parent.mkdir(parents=True, exist_ok=True)
    encoded = content.encode("utf-8")
    digest = hashlib.sha256(encoded).hexdigest()
    existed_before = target.exists()
    if existed_before:
        current = target.read_bytes()
        if current == encoded:
            return WriteResult(
                status="unchanged",
                path=str(target),
                bytes_written=0,
                sha256=digest,
            )
        if not overwrite:
            raise FileExistsError(f"target already exists: {target.name}")

    temp_name: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            newline="\n",
            prefix=".valuation-report-mcp-",
            suffix=".tmp",
            dir=target.parent,
            delete=False,
        ) as temp_file:
            temp_file.write(content)
            temp_name = temp_file.name
        os.replace(temp_name, target)
    finally:
        if temp_name and Path(temp_name).exists():
            Path(temp_name).unlink()

    return WriteResult(
        status="written" if existed_before else "created",
        path=str(target),
        bytes_written=len(encoded),
        sha256=digest,
    )


def write_report(config: WorkspaceConfig, request: DraftReportRequest) -> WriteResult:
    filename = request.output_filename or f"{request.company_name}_企业价值评估报告_初稿.md"
    requested_path = Path(filename)
    if requested_path.is_absolute() or requested_path.name != filename:
        raise ValueError("output_filename must be a filename, not a path")
    safe_filename = sanitize_markdown_filename(filename)
    target = resolve_within(config.reports, safe_filename)
    return _atomic_write(target, render_report(request), request.overwrite)


def render_company_knowledge(request: CompanyKnowledgeRequest) -> str:
    notes = _bullet_markdown(
        request.notes,
        "本信息由AI辅助收集，仅供参考，正式使用前请核对官方披露数据。",
    )
    sources = _source_markdown(request.sources)
    return f"""# {request.company_name} 企业信息档案

- **信息获取日期**：{request.information_date}
- **所属行业**：{request.industry}

## 数据来源

{sources}

## 公司概况

{request.business_description}

## 财务表现

{request.financial_performance}

## 行业对比

{request.industry_comparison}

## 备注

{notes}

- 本信息由AI辅助收集，仅供参考，正式使用前请核对官方披露数据。
"""


def write_company_knowledge(
    config: WorkspaceConfig, request: CompanyKnowledgeRequest
) -> WriteResult:
    filename = sanitize_markdown_filename(f"{request.company_name}_{request.information_date}.md")
    target = resolve_within(config.knowledge_base, Path("公司档案") / filename)
    return _atomic_write(target, render_company_knowledge(request), request.overwrite)


def check_compliance(config: WorkspaceConfig, report_relative_path: str) -> ComplianceResult:
    path = resolve_within(config.root, report_relative_path)
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"report not found: {report_relative_path}")
    if path.is_symlink():
        raise ValueError("refusing to read a report through a symlink")
    if path.suffix.lower() != ".md":
        raise ValueError("compliance check currently supports Markdown reports only")
    text = _read_text_file(path, config.max_extracted_chars)

    missing_sections = [section for section in REQUIRED_REPORT_SECTIONS if section not in text]
    detected_methods = [method for method in ALLOWED_METHODS if method in text]
    source_url_count = len(re.findall(r"https?://[^\s)>]+", text))
    has_disclaimer = any(
        marker in text for marker in ("不构成正式资产评估报告", "仅供参考", "初稿")
    )

    passed_items: list[ComplianceItem] = []
    manual_review_items: list[ComplianceItem] = []
    if not missing_sections:
        passed_items.append(
            ComplianceItem(item="报告结构", status="通过", detail="必要章节均已出现。")
        )
    else:
        manual_review_items.append(
            ComplianceItem(
                item="报告结构",
                status="需补充",
                detail=f"缺少章节：{', '.join(missing_sections)}",
            )
        )
    if detected_methods:
        passed_items.append(
            ComplianceItem(
                item="评估方法",
                status="通过",
                detail=f"检测到准则允许的方法：{', '.join(detected_methods)}",
            )
        )
    else:
        manual_review_items.append(
            ComplianceItem(
                item="评估方法",
                status="需补充",
                detail="未检测到收益法、市场法、资产基础法或成本法。",
            )
        )
    if source_url_count:
        passed_items.append(
            ComplianceItem(
                item="公开来源",
                status="通过",
                detail=f"检测到{source_url_count}个URL。",
            )
        )
    else:
        manual_review_items.append(
            ComplianceItem(
                item="公开来源",
                status="需补充",
                detail="未检测到URL；关键数据应注明可追溯来源。",
            )
        )
    if has_disclaimer:
        passed_items.append(
            ComplianceItem(item="使用限制", status="通过", detail="已检测到初稿或非正式报告声明。")
        )
    else:
        manual_review_items.append(
            ComplianceItem(
                item="使用限制",
                status="需补充",
                detail="未检测到初稿、仅供参考或非正式报告声明。",
            )
        )

    manual_review_items.extend(
        [
            ComplianceItem(
                item="业务承接",
                status="人工复核",
                detail="确认委托方、经济行为、评估目的、对象、范围、基准日和报告使用人。",
            ),
            ComplianceItem(
                item="资料与程序",
                status="人工复核",
                detail="核验申报资料、权属、现场调查、预测参数、期后事项和工作底稿。",
            ),
            ComplianceItem(
                item="最终结论",
                status="人工复核",
                detail="由资产评估师复核方法、参数、敏感性和多方法结论形成过程。",
            ),
        ]
    )
    overall = (
        "通过（初稿自检）"
        if not missing_sections and detected_methods and source_url_count and has_disclaimer
        else "需补充"
    )
    return ComplianceResult(
        report_path=str(path),
        overall_status=overall,
        passed_items=passed_items,
        manual_review_items=manual_review_items,
        missing_sections=missing_sections,
        source_url_count=source_url_count,
        detected_methods=detected_methods,
    )
